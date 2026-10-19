"""Scrive i fogli Excel dedicati a campioni, allarmi, riepilogo, legenda e parametri biometrici."""

from __future__ import annotations

import time
from typing import Any, Callable, Optional

from openpyxl.worksheet.worksheet import Worksheet

from drone.data.biometric_report_stats import (
    MAX_GAP_SEC,
    alarm_cause_label,
    format_duration,
    mean_abs_delta,
    monitored_duration_sec,
    relative_time,
    sample_hr_out_of_band,
    sample_spo2_out_of_band,
    sample_step_sec,
    state_durations,
    state_label,
    telemetry_gaps,
    valid_reading_percent,
)
from drone.data.excel_tables import (
    column_notes,
    write_data_sheet,
    write_legend_sheet as write_legend_blocks,
    write_parameters_sheet as write_parameters_block,
    write_stats_table,
)
from drone.data.flight_excel_style import (
    ALERT_FILL,
    ALERT_FONT,
    ALERT_HEADER_FILL,
    BASE_FONT,
    CELL_BORDER,
    CENTER,
    FMT_INT,
    FMT_PCT,
    FMT_S,
    FMT_TIME,
    LEFT,
    MEASURE_HEADER_FILL,
    RIGHT,
    STATUS_HEADER_FILL,
    TAB_ALARMS,
    TAB_LEGEND,
    TAB_PARAMETERS,
    TAB_WATCH,
    TARGET_HEADER_FILL,
    TIME_HEADER_FILL,
    WARN_FILL,
    WARN_FONT,
    bool_it,
    clock as as_clock,
    fmt_num_it,
    kv,
    metric_card,
    prepare_summary_sheet,
    section_title,
    table_header,
)
from drone.data.flight_report_stats import (
    column_stats,
    percent_within,
    total_duration_sec,
)

# Battito e saturazione si leggono con un decimale. Le stringhe valgono anche
# per autofit_columns, che riconosce il formato a un decimale.
FMT_BPM = "0.0"
FMT_SPO2 = "0.0"
FMT_HZ = "0.0"
FMT_SEC = "0.0"

_STAT_LABELS = {
    "hr_grezzo": "Battito misurato [bpm]",
    "hr_filtrato": "Battito filtrato [bpm]",
    "spo2_grezzo": "Saturazione SpO₂ misurata [%]",
    "spo2_filtrato": "Saturazione SpO₂ filtrata [%]",
}

_STAT_COLUMNS = ["hr_grezzo", "hr_filtrato", "spo2_grezzo", "spo2_filtrato"]

SAMPLES_SHEET = "Orologio"


def _delta(filtered: Any, raw: Any) -> Optional[float]:
    if filtered is None or raw is None:
        return None
    return float(filtered) - float(raw)


def _sample_columns(
    thresholds, clock: tuple[str, float]
) -> list[tuple[str, Callable, Optional[str]]]:
    return [
        ("Orario", lambda c: as_clock(c["timestamp"]), FMT_TIME),
        ("Tempo dall'avvio [s]", lambda c: relative_time(c, clock), FMT_S),
        ("Stato dell'orologio", lambda c: state_label(c.get("stato")), None),
        ("Lettura valida", lambda c: bool_it(c.get("lettura_valida")), None),
        ("Battito misurato [bpm]", lambda c: c.get("hr_grezzo"), FMT_BPM),
        ("Battito filtrato [bpm]", lambda c: c.get("hr_filtrato"), FMT_BPM),
        ("Differenza battito (filtrato − misurato) [bpm]",
         lambda c: _delta(c.get("hr_filtrato"), c.get("hr_grezzo")), FMT_BPM),
        ("Saturazione SpO₂ misurata [%]", lambda c: c.get("spo2_grezzo"), FMT_SPO2),
        ("Saturazione SpO₂ filtrata [%]", lambda c: c.get("spo2_filtrato"), FMT_SPO2),
        ("Differenza SpO₂ (filtrata − misurata) [%]",
         lambda c: _delta(c.get("spo2_filtrato"), c.get("spo2_grezzo")), FMT_SPO2),
        ("Battito fuori soglia",
         lambda c: bool_it(sample_hr_out_of_band(c, thresholds)), None),
        ("SpO₂ sotto soglia",
         lambda c: bool_it(sample_spo2_out_of_band(c, thresholds)), None),
    ]


def _alarm_columns(
    thresholds, clock: tuple[str, float]
) -> list[tuple[str, Callable, Optional[str]]]:
    return [
        ("Orario", lambda e: as_clock(e["timestamp"]), FMT_TIME),
        ("Tempo dall'avvio [s]", lambda e: relative_time(e, clock), FMT_S),
        ("Battito [bpm]", lambda e: e.get("bpm"), FMT_BPM),
        ("Saturazione SpO₂ [%]", lambda e: e.get("spo2"), FMT_SPO2),
        ("Causa", lambda e: alarm_cause_label(e, thresholds), None),
    ]


def write_samples_sheet(ws: Worksheet, samples, thresholds, clock) -> None:
    """Scrive il foglio con i campioni biometrici della sessione."""
    time_cols = {"Orario", "Tempo dall'avvio [s]"}
    status_cols = {"Stato dell'orologio", "Lettura valida"}
    hr_cols = {"Battito misurato [bpm]", "Battito filtrato [bpm]", "Differenza battito (filtrato − misurato) [bpm]"}
    spo2_cols = {
        "Saturazione SpO₂ misurata [%]", "Saturazione SpO₂ filtrata [%]", "Differenza SpO₂ (filtrata − misurata) [%]"
    }
    alert_cols = {"Battito fuori soglia", "SpO₂ sotto soglia"}
    header_fills = {}
    header_fills.update({h: TIME_HEADER_FILL for h in time_cols})
    header_fills.update({h: STATUS_HEADER_FILL for h in status_cols})
    header_fills.update({h: MEASURE_HEADER_FILL for h in hr_cols})
    header_fills.update({h: TARGET_HEADER_FILL for h in spo2_cols})
    header_fills.update({h: ALERT_HEADER_FILL for h in alert_cols})

    write_data_sheet(
        ws,
        samples,
        _sample_columns(thresholds, clock),
        highlights={
            "Stato dell'orologio": ("Allarme", ALERT_FILL, ALERT_FONT),
            "Battito fuori soglia": ("Sì", ALERT_FILL, ALERT_FONT),
            "SpO₂ sotto soglia": ("Sì", ALERT_FILL, ALERT_FONT),
            "Lettura valida": ("No", WARN_FILL, WARN_FONT),
        },
        notes=column_notes(_LEGEND_SAMPLES),
        title="Monitoraggio smartwatch — campioni biometrici",
        subtitle=(
            "Un campione per ogni acquisizione ricevuta: stato del sensore, battito, "
            "saturazione SpO₂, valori filtrati e superamento delle soglie configurate."
        ),
        tab_color=TAB_WATCH,
        header_fills=header_fills,
        freeze_columns=4,
        center_headers={
            "Stato dell'orologio", "Lettura valida",
            "Battito fuori soglia", "SpO₂ sotto soglia",
        },
        table_name="Tbl_Orologio",
    )


def write_alarms_sheet(ws: Worksheet, events, thresholds, clock) -> None:
    """Scrive il foglio con gli eventi di allarme biometrici."""
    header_fills = {
        "Orario": TIME_HEADER_FILL,
        "Tempo dall'avvio [s]": TIME_HEADER_FILL,
        "Battito [bpm]": MEASURE_HEADER_FILL,
        "Saturazione SpO₂ [%]": TARGET_HEADER_FILL,
        "Causa": ALERT_HEADER_FILL,
    }
    write_data_sheet(
        ws, events, _alarm_columns(thresholds, clock),
        notes=column_notes(_LEGEND_ALARMS),
        title="Allarmi registrati",
        subtitle=(
            "Eventi di allarme generati durante la sessione, con valori rilevati "
            "e causa ricavata dalle soglie configurate."
        ),
        tab_color=TAB_ALARMS,
        header_fills=header_fills,
        freeze_columns=2,
        table_name="Tbl_Allarmi",
    )


def _write_overview(ws: Worksheet, row: int, samples, titolo: str) -> int:
    row = section_title(ws, row, titolo, span=6)
    row = kv(ws, row, "Campioni biometrici ricevuti", len(samples), FMT_INT)
    row = kv(ws, row, "Durata della sessione [s]", total_duration_sec(samples), FMT_SEC)
    row = kv(
        ws, row, "Durata con telemetria disponibile [s]",
        monitored_duration_sec(samples), FMT_SEC,
    )
    # La cadenza è la mediana dei passi, non i campioni diviso la durata: con
    # un buco nella telemetria la media direbbe che l'orologio pubblica piano.
    passo = sample_step_sec(samples)
    row = kv(ws, row, "Intervallo tipico tra campioni [s]", passo, FMT_SEC)
    if passo > 0:
        row = kv(ws, row, "Frequenza di acquisizione [campioni/s]", 1.0 / passo, FMT_HZ)
    t0 = float(samples[0]["timestamp"])
    t1 = float(samples[-1]["timestamp"])
    row = kv(
        ws, row, "Intervallo orario della sessione",
        f"da {time.strftime('%H:%M:%S', time.localtime(t0))} "
        f"a {time.strftime('%H:%M:%S', time.localtime(t1))}",
    )
    valide = valid_reading_percent(samples)
    if valide is not None:
        row = kv(ws, row, "Campioni con misura valida [%]", valide, FMT_PCT)
    return row + 1


def _write_outcome(ws: Worksheet, row: int, samples, events, thresholds, titolo: str) -> int:
    row = section_title(ws, row, titolo, span=6)
    row = kv(ws, row, "Allarmi biometrici registrati", len(events), FMT_INT)
    row = kv(
        ws, row, "Campioni con battito fuori soglia",
        sum(1 for c in samples if sample_hr_out_of_band(c, thresholds)),
        FMT_INT,
    )
    row = kv(
        ws, row, "Campioni con SpO₂ sotto soglia",
        sum(1 for c in samples if sample_spo2_out_of_band(c, thresholds)),
        FMT_INT,
    )

    entro_hr = percent_within(
        samples,
        lambda c: not sample_hr_out_of_band(c, thresholds),
        lambda c: c.get("bpm") is not None,
    )
    if entro_hr is not None:
        row = kv(
            ws, row,
            f"Percentuale di campioni con battito nei limiti ({thresholds.bpm_min_in}–{thresholds.bpm_max_in} bpm)",
            entro_hr, FMT_PCT,
        )
    entro_spo2 = percent_within(
        samples,
        lambda c: not sample_spo2_out_of_band(c, thresholds),
        lambda c: c.get("spo2") is not None,
    )
    if entro_spo2 is not None:
        row = kv(
            ws, row,
            f"Percentuale di campioni con SpO₂ nei limiti (≥ {thresholds.spo2_min_in}%)",
            entro_spo2, FMT_PCT,
        )
    return row + 1


def _write_state_table(ws: Worksheet, row: int, samples, titolo: str) -> int:
    durate = state_durations(samples)
    if not durate:
        return row
    totale = sum(durata for _codice, durata in durate)
    row = section_title(ws, row, titolo, span=6)
    row = table_header(ws, row, ["Stato dell'orologio", "Durata [s]", "Quota della durata monitorata [%]", "Durata leggibile"])
    for codice, durata in durate:
        nome = ws.cell(row=row, column=1, value=state_label(codice))
        nome.font = BASE_FONT
        nome.alignment = LEFT
        nome.border = CELL_BORDER
        secondi = ws.cell(row=row, column=2, value=durata)
        secondi.font = BASE_FONT
        secondi.number_format = FMT_SEC
        secondi.alignment = RIGHT
        secondi.border = CELL_BORDER
        quota = ws.cell(
            row=row, column=3, value=(100.0 * durata / totale) if totale else 0.0
        )
        quota.font = BASE_FONT
        quota.number_format = FMT_PCT
        quota.alignment = RIGHT
        quota.border = CELL_BORDER
        leggibile = ws.cell(row=row, column=4, value=format_duration(durata))
        leggibile.font = BASE_FONT
        leggibile.alignment = CENTER
        leggibile.border = CELL_BORDER
        row += 1
    return row + 1


def _write_filter_section(ws: Worksheet, row: int, samples, titolo: str) -> int:
    scarto_hr = mean_abs_delta(samples, "hr_grezzo", "hr_filtrato")
    scarto_spo2 = mean_abs_delta(samples, "spo2_grezzo", "spo2_filtrato")
    if scarto_hr is None and scarto_spo2 is None:
        return row
    row = section_title(ws, row, titolo, span=6)
    if scarto_hr is not None:
        row = kv(
            ws, row, "Correzione media assoluta del filtro sul battito [bpm]",
            scarto_hr, FMT_BPM,
        )
    if scarto_spo2 is not None:
        row = kv(
            ws, row, "Correzione media assoluta del filtro sulla SpO₂ [%]",
            scarto_spo2, FMT_SPO2,
        )
    return row + 1


def _write_gaps_section(ws: Worksheet, row: int, samples, titolo: str) -> int:
    buchi = telemetry_gaps(samples)
    if not buchi:
        return row
    row = section_title(ws, row, titolo, span=6)
    row = kv(ws, row, "Interruzioni della telemetria [n.]", len(buchi), FMT_INT)
    row = kv(
        ws, row, "Durata complessiva delle interruzioni [s]",
        sum(buco["duration"] for buco in buchi), FMT_SEC,
    )
    for buco in buchi:
        cella = ws.cell(
            row=row, column=1,
            value=(
                f"• da {fmt_num_it(buco['start'], 1)} s a {fmt_num_it(buco['end'], 1)} s "
                f"({format_duration(buco['duration'])} senza dati)"
            ),
        )
        cella.font = BASE_FONT
        cella.alignment = LEFT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        row += 1
    return row + 1


def write_summary_sheet(ws: Worksheet, logger, thresholds) -> None:
    """Scrive il foglio riepilogativo del report."""
    samples = logger.samples()
    events = logger.events()

    row = prepare_summary_sheet(
        ws,
        "Riepilogo sessione biometrica",
        f"Generato il {time.strftime('%d/%m/%Y alle %H:%M:%S')}",
    )

    if samples:
        metric_card(
            ws, row, 1, "Durata sessione [s]",
            total_duration_sec(samples), number_format=FMT_SEC,
        )
        metric_card(ws, row, 3, "Campioni biometrici ricevuti", len(samples), number_format=FMT_INT)
        metric_card(
            ws, row, 5, "Allarmi registrati", len(events), number_format=FMT_INT,
            fill=ALERT_FILL if events else None,
        )
        row += 3

        row = _write_overview(ws, row, samples, "Panoramica")
        row = _write_outcome(
            ws, row, samples, events, thresholds, "Soglie e allarmi"
        )
        row = _write_state_table(ws, row, samples, "Durata per stato dell'orologio")

        stats = column_stats(samples, _STAT_COLUMNS)
        if stats:
            row = section_title(ws, row, "Statistiche dei valori biometrici", span=6)
            row = write_stats_table(
                ws, row, stats,
                labels=_STAT_LABELS,
                number_format=lambda nome: FMT_SPO2 if "spo2" in nome else FMT_BPM,
            )
            row += 1
        row = _write_filter_section(
            ws, row, samples, "Effetto del filtro sui valori biometrici"
        )
        _write_gaps_section(ws, row, samples, "Interruzioni della telemetria")
    else:
        metric_card(ws, row, 1, "Sessione", "Nessun dato")


# Cosa contiene ciascun foglio, nell'ordine in cui compaiono nel file.
_LEGEND_SHEETS = {
    "Riepilogo": "Il riassunto della sessione: durata, qualità della telemetria, allarmi, battito e saturazione SpO₂.",
    "Parametri": "Le soglie di allarme per battito e saturazione SpO₂ e i parametri della telemetria smartwatch.",
    SAMPLES_SHEET: "Tutti i campioni inviati dallo smartwatch: stato del sensore, battito, saturazione SpO₂ e valori filtrati.",
    "Allarmi": "Gli eventi di allarme biometrico registrati durante la sessione, con valori e causa.",
}

# Parole che compaiono nei nomi delle colonne e dei parametri.
_LEGEND_WORDS = [
    ("bpm", "Battiti del cuore al minuto."),
    ("Saturazione SpO₂ [%]", "Percentuale di emoglobina ossigenata stimata dal sensore; è la misura indicata come SpO₂."),
    ("Misurato / filtrato", "«Misurato» è il valore fornito dal sensore; «filtrato» è il valore dopo la riduzione degli sbalzi e del rumore di misura."),
    ("Limiti", "I valori oltre i quali l'orologio dà l'allarme: sono scritti nel foglio «Parametri»."),
]

_LEGEND_SUMMARY = [
    ("Campioni biometrici ricevuti", f"Numero totale di campioni biometrici ricevuti e salvati nel foglio «{SAMPLES_SHEET}»."),
    ("Durata della sessione [s]", "Intervallo di tempo tra il primo e l'ultimo campione registrato nella sessione."),
    ("Durata con telemetria disponibile [s]", "Durata effettivamente coperta dai campioni ricevuti, escludendo le interruzioni della telemetria più lunghe della soglia configurata."),
    ("Intervallo tipico tra campioni [s]", "Intervallo tipico tra due campioni consecutivi, calcolato come mediana e senza considerare le interruzioni lunghe."),
    ("Frequenza di acquisizione [campioni/s]", "Frequenza tipica di acquisizione, espressa in campioni ricevuti al secondo."),
    ("Intervallo orario della sessione", "Ora del primo e dell'ultimo campione registrato nella sessione."),
    ("Campioni con misura valida [%]", "Percentuale dei campioni in cui il sensore ha dichiarato valida la misura biometrica."),
    ("Allarmi biometrici registrati", "Numero di eventi di allarme biometrico generati per battito e/o saturazione SpO₂."),
    ("Campioni con battito fuori soglia", "Numero di campioni in cui il battito era inferiore o superiore alle soglie di ingresso dell'allarme."),
    ("Campioni con SpO₂ sotto soglia", "Numero di campioni in cui la saturazione SpO₂ era inferiore alla soglia di ingresso dell'allarme."),
    ("Percentuale di campioni con battito nei limiti (…)", "Percentuale dei campioni con battito disponibile che rientrano nelle soglie indicate."),
    ("Percentuale di campioni con SpO₂ nei limiti (≥ …%)", "Percentuale dei campioni con SpO₂ disponibile che rispettano la soglia minima indicata."),
    ("Durata per stato dell'orologio", "Ripartizione della durata monitorata tra i diversi stati operativi dello smartwatch."),
    ("Durata [s]", "Secondi passati in quello stato."),
    ("Quota della durata monitorata [%]", "Percentuale della durata coperta dalla telemetria trascorsa in quello stato."),
    ("Durata leggibile", "La stessa durata espressa in forma più leggibile, usando minuti e secondi."),
    ("Grandezza", "Parametro biometrico per cui vengono riportate le statistiche."),
    ("Minimo / Massimo", "Il valore più basso e quello più alto della colonna."),
    ("Media", "Il valore medio della colonna."),
    ("Deviazione standard", "Misura quanto i valori sono dispersi attorno alla media: più è bassa, più il parametro è stabile. «—» quando non è calcolabile."),
    ("Correzione media assoluta del filtro sul battito [bpm]", "Media del valore assoluto della correzione introdotta dal filtro sul battito."),
    ("Correzione media assoluta del filtro sulla SpO₂ [%]", "Media del valore assoluto della correzione introdotta dal filtro sulla saturazione SpO₂."),
    ("Interruzioni della telemetria [n.]", f"Numero di interruzioni in cui non sono arrivati campioni per più di {fmt_num_it(MAX_GAP_SEC, 0)} secondi."),
    ("Durata complessiva delle interruzioni [s]", "Somma della durata di tutte le interruzioni della telemetria rilevate nella sessione."),
    ("• da … s a … s", "Intervallo temporale di una singola interruzione, espresso in secondi dall'inizio della sessione."),
]

_LEGEND_PARAMETERS = [
    ("Allarme se il battito scende sotto [bpm]", "Se il battito scende sotto questo valore, l'orologio controlla per qualche secondo e poi dà l'allarme."),
    ("Allarme se il battito sale sopra [bpm]", "Se il battito sale sopra questo valore, l'orologio controlla per qualche secondo e poi dà l'allarme."),
    ("Fine allarme quando il battito risale fino a [bpm]", "Dopo un allarme per battito basso, l'allarme finisce quando il battito risale fino a questo valore."),
    ("Fine allarme quando il battito riscende fino a [bpm]", "Dopo un allarme per battito alto, l'allarme finisce quando il battito riscende fino a questo valore."),
    ("Allarme se la saturazione SpO₂ scende sotto [%]", "Se la saturazione SpO₂ scende sotto questo valore, lo smartwatch verifica la condizione e può generare l'allarme."),
    ("Fine allarme quando la saturazione SpO₂ risale fino a [%]", "Dopo un allarme per SpO₂ bassa, l'allarme termina quando la saturazione risale fino a questo valore."),
    ("Intervallo nominale tra campioni", "Intervallo di pubblicazione previsto dal firmware dello smartwatch."),
    ("Intervallo tipico osservato tra campioni", "Intervallo mediano effettivamente osservato tra campioni consecutivi, ignorando le interruzioni lunghe."),
    ("Campioni biometrici ricevuti", "Numero totale di campioni ricevuti dallo smartwatch durante la sessione."),
    ("Soglia per rilevare un’interruzione", "Se tra due campioni passa più di questo tempo, l'intervallo viene classificato come interruzione della telemetria."),
    ("Broker MQTT di ricezione", "Indirizzo e porta del broker MQTT attraverso cui vengono ricevuti i campioni dello smartwatch."),
]

_LEGEND_SAMPLES = [
    ("Orario", "Ora del computer in cui è stato ricevuto il campione."),
    ("Tempo dall'avvio [s]", "Secondi trascorsi dall'inizio della sessione, usando lo stesso riferimento temporale del foglio «Allarmi»."),
    ("Stato dell'orologio", "Cosa stava facendo l'orologio (rosso se in allarme). Il significato di ogni valore è nella Legenda."),
    ("Lettura valida", "Sì se il sensore ha prodotto una misura biometrica valida; No (giallo) se battito e SpO₂ del campione non sono affidabili."),
    ("Battito misurato [bpm]", "Frequenza cardiaca fornita dal sensore prima del filtraggio. Vuoto quando non è disponibile una misura."),
    ("Battito filtrato [bpm]", "Frequenza cardiaca dopo il filtraggio; è il valore stabilizzato usato dalla logica di monitoraggio."),
    ("Differenza battito (filtrato − misurato) [bpm]", "Correzione con segno introdotta dal filtro: valore filtrato meno valore misurato."),
    ("Saturazione SpO₂ misurata [%]", "Valore di saturazione SpO₂ fornito dal sensore prima del filtraggio. Vuoto quando non è disponibile una misura."),
    ("Saturazione SpO₂ filtrata [%]", "Saturazione SpO₂ dopo il filtraggio; è il valore stabilizzato usato dalla logica di monitoraggio."),
    ("Differenza SpO₂ (filtrata − misurata) [%]", "Correzione con segno introdotta dal filtro: SpO₂ filtrata meno SpO₂ misurata."),
    ("Battito fuori soglia", "Sì (rosso) quando il battito è sotto la soglia minima o sopra la soglia massima dell'allarme."),
    ("SpO₂ sotto soglia", "Sì (rosso) quando la saturazione SpO₂ è inferiore alla soglia minima dell'allarme."),
]

# Stessi nomi di STATE_LABELS, nello stesso ordine.
_LEGEND_WATCH_STATES = [
    ("Missione non avviata", "La missione non è ancora iniziata: l'orologio non controlla ancora battito e SpO₂."),
    ("Ricerca segnale", "L'orologio sta cercando di leggere battito e SpO₂ (sensore appoggiato male o segnale debole)."),
    ("Normale", "Battito e ossigeno sono a posto."),
    ("Verifica", "Un valore è fuori dai limiti: l'orologio aspetta qualche secondo per essere sicuro, prima di dare l'allarme."),
    ("Allarme", "L'orologio sta dando l'allarme: battito o ossigeno fuori dai limiti, oppure un avviso arrivato dal drone (dispositivi di protezione mancanti o ingresso in una zona vietata)."),
    ("Silenziato", "L'operatore ha spento l'allarme con il pulsante; se il problema continua, l'allarme riparte dopo un po'."),
    ("Guasto del sensore", "Il sensore di battito e SpO₂ non funziona."),
]

_LEGEND_ALARMS = [
    ("Orario", "Ora del computer in cui è arrivato l'allarme."),
    ("Tempo dall'avvio [s]", f"Secondi passati dal primo dato della sessione, contati come nel foglio «{SAMPLES_SHEET}»."),
    ("Battito [bpm]", "Il battito al momento dell'allarme."),
    ("Saturazione SpO₂ [%]", "Saturazione SpO₂ disponibile al momento dell'allarme."),
    ("Causa", "Cosa ha fatto scattare l'allarme: il battito, la SpO₂ o entrambi."),
]


def write_legend_sheet(
    ws: Worksheet,
    *,
    include_parameters: bool = False,
    include_samples: bool,
    include_alarms: bool,
) -> None:
    """Scrive il foglio che descrive campi, unità e significato delle colonne."""
    fogli = ["Riepilogo"]
    if include_parameters:
        fogli.append("Parametri")
    if include_samples:
        fogli.append(SAMPLES_SHEET)
    if include_alarms:
        fogli.append("Allarmi")

    blocks = [
        ("Fogli del file", [(f"«{nome}»", _LEGEND_SHEETS[nome]) for nome in fogli]),
        ("Parole usate nel file", _LEGEND_WORDS),
    ]
    # Senza campioni il riepilogo resta vuoto: non c'è niente da spiegare.
    if include_samples:
        blocks.append(("Foglio «Riepilogo»", _LEGEND_SUMMARY))
    if include_parameters:
        blocks.append(("Foglio «Parametri»", _LEGEND_PARAMETERS))
    if include_samples:
        blocks.append((f"Foglio «{SAMPLES_SHEET}»", _LEGEND_SAMPLES))
        blocks.append(("Cosa vuol dire ogni valore di «Stato dell'orologio»", _LEGEND_WATCH_STATES))
    if include_alarms:
        blocks.append(("Foglio «Allarmi»", _LEGEND_ALARMS))
    write_legend_blocks(ws, blocks, name_width=44, meaning_width=86, tab_color=TAB_LEGEND)


def write_parameters_sheet(ws: Worksheet, parameters: list[tuple[str, list[tuple[str, str]]]]) -> None:
    """Scrive il foglio con i parametri di configurazione della sessione."""
    write_parameters_block(
        ws,
        parameters,
        subtitle="I limiti dell'allarme e il collegamento con cui è stata registrata la sessione.",
        note=(
            "I limiti e la decisione di dare l'allarme stanno nel programma "
            "dell'orologio: qui sono scritti solo per leggere i dati."
        ),
        tab_color=TAB_PARAMETERS,
    )
