"""Costruisce il workbook Excel riepilogativo di una sessione biometrica."""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from openpyxl import Workbook

from drone.data.biometric_excel_sheets import (
    SAMPLES_SHEET,
    write_alarms_sheet,
    write_legend_sheet,
    write_parameters_sheet,
    write_samples_sheet,
    write_summary_sheet,
)
from drone.data.biometric_report_stats import (
    MAX_GAP_SEC,
    NOMINAL_STEP_SEC,
    sample_step_sec,
    session_clock,
)
from drone.data.flight_excel_style import fmt_num_it


def collect_session_parameters(
    logger, app_config
) -> list[tuple[str, list[tuple[str, str]]]]:
    """Raccoglie i parametri da riportare nel report della sessione biometrica."""
    soglie = app_config.smartwatch_thresholds
    samples = logger.samples()

    telemetria = [
        ("Intervallo nominale tra campioni", f"{fmt_num_it(NOMINAL_STEP_SEC, 1)} s"),
    ]
    if samples:
        telemetria.append(
            ("Intervallo tipico osservato tra campioni", f"{fmt_num_it(sample_step_sec(samples), 2)} s")
        )
        telemetria.append(("Campioni biometrici ricevuti", str(len(samples))))
    telemetria.append(
        ("Soglia per rilevare un’interruzione", f"{fmt_num_it(MAX_GAP_SEC, 1)} s")
    )
    telemetria.append(
        ("Broker MQTT di ricezione", f"{app_config.mqtt_broker_ip}:{app_config.mqtt_broker_port}")
    )

    return [
        ("Limiti del battito", [
            ("Allarme se il battito scende sotto [bpm]", str(soglie.bpm_min_in)),
            ("Allarme se il battito sale sopra [bpm]", str(soglie.bpm_max_in)),
            ("Fine allarme quando il battito risale fino a [bpm]", str(soglie.bpm_min_out)),
            ("Fine allarme quando il battito riscende fino a [bpm]", str(soglie.bpm_max_out)),
        ]),
        ("Limiti della saturazione SpO₂", [
            ("Allarme se la saturazione SpO₂ scende sotto [%]", str(soglie.spo2_min_in)),
            ("Fine allarme quando la saturazione SpO₂ risale fino a [%]", str(soglie.spo2_min_out)),
        ]),
        ("Telemetria smartwatch", telemetria),
    ]


def _build_workbook(
    logger,
    thresholds,
    parameters: Optional[list[tuple[str, list[tuple[str, str]]]]] = None,
) -> Workbook:
    samples = logger.samples()
    allarmi = logger.events()
    clock = session_clock(samples, allarmi)

    wb = Workbook()
    wb.properties.creator = "Work-Guardian"
    wb.properties.title = "Report sessione biometrica"
    wb.properties.subject = "Dati e riepilogo della sessione biometrica"
    wb.properties.description = "Report Work-Guardian con riepilogo, parametri, dati smartwatch, allarmi e legenda."
    summary_ws = wb.active
    summary_ws.title = "Riepilogo"
    write_summary_sheet(summary_ws, logger, thresholds)

    if parameters:
        write_parameters_sheet(wb.create_sheet("Parametri"), parameters)

    if samples:
        write_samples_sheet(
            wb.create_sheet(SAMPLES_SHEET), samples, thresholds, clock
        )

    if allarmi:
        write_alarms_sheet(wb.create_sheet("Allarmi"), allarmi, thresholds, clock)

    write_legend_sheet(
        wb.create_sheet("Legenda"),
        include_parameters=bool(parameters),
        include_samples=bool(samples),
        include_alarms=bool(allarmi),
    )

    return wb


def save_session_workbook(
    logger, output_path: Path, thresholds, parameters=None
) -> Path:
    """Genera e salva il workbook Excel della sessione."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = _build_workbook(logger, thresholds, parameters)
    wb.save(output_path)
    return output_path
