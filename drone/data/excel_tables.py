"""Utility comuni per creare tabelle, note e fogli di supporto nei report Excel."""

from __future__ import annotations

import math
import re
import textwrap
from typing import Any, Callable, Optional

from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from drone.data.flight_excel_style import (
    BASE_FONT,
    CELL_BORDER,
    CENTER,
    HEADER_FILL,
    HEADER_FONT,
    LABEL_FONT,
    LEFT,
    RIGHT,
    SOFT_BORDER,
    SUBTITLE_FONT,
    TITLE_FONT,
    WRAP_CENTER,
    WRAP_LEFT,
    ZEBRA_FILL,
    add_back_link,
    autofit_columns,
    configure_print,
    report_title,
    section_title,
    setup_sheet,
    table_header,
)

# I fogli dei report di volo e biometrici sono costruiti con gli stessi
# mattoni: tabella campioni, statistiche, parametri e legenda. Qui sono raccolte
# le regole visive comuni, così i due Excel restano coerenti.

Column = tuple[str, Callable[[Any], Any], Optional[str]]
Stat = tuple[str, float, float, float, Optional[float]]

_NOTE_AUTHOR = "Work-Guardian"
_NOTE_WIDTH_PX = 340
_NOTE_CHARS_PER_LINE = 50
_NOTE_LINE_PX = 15
_NOTE_PADDING_PX = 12

# Altezza di una riga di testo dell'intestazione (Arial 10 grassetto) e
# altezza minima della riga delle intestazioni, in punti.
HEADER_LINE_PT = 13
_HEADER_MIN_HEIGHT_PT = 38


def column_notes(pairs: list[tuple[str, str]]) -> dict[str, str]:
    """Costruisce le note descrittive associate alle colonne di un foglio dati."""
    notes: dict[str, str] = {}
    for name, meaning in pairs:
        if "X / Y / Z" in name:
            for axis in "XYZ":
                notes[name.replace("X / Y / Z", axis)] = meaning
        else:
            notes[name] = meaning
    return notes


def _header_note(text: str) -> Comment:
    lines = sum(
        max(1, math.ceil(len(paragraph) / _NOTE_CHARS_PER_LINE))
        for paragraph in text.split("\n")
    )
    return Comment(
        text,
        _NOTE_AUTHOR,
        height=lines * _NOTE_LINE_PX + _NOTE_PADDING_PX,
        width=_NOTE_WIDTH_PX,
    )


def _safe_table_name(name: str) -> str:
    clean = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if not clean or clean[0].isdigit():
        clean = f"Data_{clean}"
    return f"Tbl_{clean}"[:240]


def write_data_sheet(
    ws: Worksheet,
    entries,
    columns: list[Column],
    *,
    highlights: Optional[dict[str, tuple[str, PatternFill, Font]]] = None,
    notes: Optional[dict[str, str]] = None,
    title: Optional[str] = None,
    subtitle: Optional[str] = None,
    tab_color: Optional[str] = None,
    header_fills: Optional[dict[str, PatternFill]] = None,
    freeze_columns: int = 2,
    center_headers: Optional[set[str]] = None,
    table_name: Optional[str] = None,
) -> None:
    """Scrive un foglio dati pensato per essere letto, filtrato e navigato.

    La riga 5 contiene le intestazioni; sopra ci sono titolo, spiegazione e link
    al riepilogo. Le prime colonne possono restare bloccate durante lo scroll.
    """
    rows = list(entries)
    headers = [header for header, _getter, _fmt in columns]
    last_col_index = len(columns)
    last_col = get_column_letter(last_col_index)

    setup_sheet(ws, tab_color=tab_color, zoom=85)
    configure_print(
        ws, orientation="landscape",
        fit_to_width=2 if last_col_index > 14 else 1,
        repeat_rows="$5:$5",
    )

    report_title(
        ws,
        title or ws.title,
        subtitle or "Dati registrati durante la sessione.",
        span=last_col_index,
    )
    add_back_link(ws, row=3)

    header_row = 5
    data_start_row = 6

    for col_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=header)
        cell.font = HEADER_FONT
        cell.fill = (header_fills or {}).get(header, HEADER_FILL)
        cell.alignment = WRAP_CENTER
        cell.border = CELL_BORDER
        if notes and header in notes:
            cell.comment = _header_note(notes[header])
    ws.row_dimensions[header_row].height = 38

    center_headers = center_headers or set()
    for row_index, entry in enumerate(rows, start=data_start_row):
        for col_index, (header, getter, fmt) in enumerate(columns, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=getter(entry))
            cell.font = BASE_FONT
            cell.border = SOFT_BORDER
            if fmt is not None:
                cell.number_format = fmt
                cell.alignment = RIGHT
            elif header in center_headers:
                cell.alignment = CENTER
            else:
                cell.alignment = LEFT

        for header, (trigger_word, fill, font) in (highlights or {}).items():
            hi_cell = ws.cell(row=row_index, column=headers.index(header) + 1)
            if str(hi_cell.value) == trigger_word:
                hi_cell.fill = fill
                hi_cell.font = font
                hi_cell.alignment = CENTER

    # Una vera tabella Excel migliora filtri, selezione e righe alternate.
    last_row = data_start_row + len(rows) - 1
    if rows:
        ref = f"A{header_row}:{last_col}{last_row}"
        table = Table(
            displayName=table_name or _safe_table_name(ws.title),
            ref=ref,
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
        ws.print_area = f"A1:{last_col}{last_row}"

    # Blocca titolo/intestazione e le colonne temporali principali.
    freeze_col = get_column_letter(max(1, freeze_columns) + 1)
    ws.freeze_panes = f"{freeze_col}{data_start_row}"

    autofit_columns(ws, max_width=30)
    for col_index, (header, _getter, fmt) in enumerate(columns, start=1):
        letter = get_column_letter(col_index)
        current = ws.column_dimensions[letter].width or 10
        if header == "Orario":
            ws.column_dimensions[letter].width = 11
        elif header.startswith("Tempo dall'avvio"):
            ws.column_dimensions[letter].width = 17
        elif fmt is not None:
            ws.column_dimensions[letter].width = min(max(current, 11), 18)
        else:
            ws.column_dimensions[letter].width = min(max(current, 12), 28)

    # Le intestazioni vanno a capo nella larghezza della loro colonna: la riga
    # si alza quanto serve alla più lunga, altrimenti l'ultima riga di testo
    # verrebbe tagliata.
    header_lines = max(
        len(textwrap.wrap(
            header,
            int(ws.column_dimensions[get_column_letter(col_index)].width),
            break_long_words=False,
        ))
        for col_index, header in enumerate(headers, start=1)
    ) if headers else 1
    ws.row_dimensions[header_row].height = max(
        _HEADER_MIN_HEIGHT_PT, header_lines * HEADER_LINE_PT + 6
    )


def write_stats_table(
    ws: Worksheet,
    row: int,
    stats: list[Stat],
    *,
    labels: dict[str, str],
    number_format: Callable[[str], str],
) -> int:
    """Scrive nel foglio una tabella compatta di statistiche descrittive."""
    row = table_header(ws, row, ["Grandezza", "Minimo", "Massimo", "Media", "Deviazione standard"])
    for item_index, (name, mn, mx, mean, std) in enumerate(stats):
        fmt = number_format(name)
        label_cell = ws.cell(row=row, column=1, value=labels.get(name, name))
        label_cell.font = BASE_FONT
        label_cell.alignment = WRAP_LEFT
        label_cell.border = CELL_BORDER
        if item_index % 2:
            label_cell.fill = ZEBRA_FILL
        for col_index, val in ((2, mn), (3, mx), (4, mean), (5, std)):
            if val is None:
                c = ws.cell(row=row, column=col_index, value="—")
                c.alignment = CENTER
            else:
                c = ws.cell(row=row, column=col_index, value=val)
                c.number_format = fmt
                c.alignment = RIGHT
            c.font = BASE_FONT
            c.border = CELL_BORDER
            if item_index % 2:
                c.fill = ZEBRA_FILL
        if len(str(labels.get(name, name))) > 24:
            ws.row_dimensions[row].height = 30
        row += 1
    return row


def write_parameters_sheet(
    ws: Worksheet,
    parameters: list[tuple[str, list[tuple[str, str]]]],
    *,
    subtitle: str,
    note: Optional[str] = None,
    tab_color: Optional[str] = None,
) -> None:
    """Scrive il foglio con i parametri di configurazione della sessione."""
    setup_sheet(ws, tab_color=tab_color, zoom=95)
    configure_print(ws, orientation="portrait", fit_to_width=1)
    report_title(ws, "Parametri della sessione", subtitle, span=2)
    add_back_link(ws, row=3)
    row = 5

    for group_title, items in parameters:
        row = section_title(ws, row, group_title, span=2)
        for item_index, (label, value) in enumerate(items):
            label_cell = ws.cell(row=row, column=1, value=label)
            label_cell.font = LABEL_FONT
            label_cell.alignment = WRAP_LEFT
            label_cell.border = CELL_BORDER
            value_cell = ws.cell(row=row, column=2, value=value)
            value_cell.font = BASE_FONT
            value_cell.alignment = WRAP_LEFT
            value_cell.border = CELL_BORDER
            if item_index % 2:
                label_cell.fill = ZEBRA_FILL
                value_cell.fill = ZEBRA_FILL
            row += 1
        row += 1

    if note:
        note_cell = ws.cell(row=row, column=1, value=note)
        note_cell.font = SUBTITLE_FONT
        note_cell.alignment = WRAP_LEFT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = 30

    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 38
    ws.freeze_panes = "A5"


def write_legend_sheet(
    ws: Worksheet,
    blocks: list[tuple[str, list[tuple[str, str]]]],
    *,
    name_width: int = 32,
    meaning_width: int = 78,
    tab_color: Optional[str] = None,
) -> None:
    """Scrive il foglio che descrive campi, unità e significato delle colonne."""
    setup_sheet(ws, tab_color=tab_color, zoom=90)
    configure_print(ws, orientation="portrait", fit_to_width=1)
    ws.column_dimensions["A"].width = name_width
    ws.column_dimensions["B"].width = meaning_width

    report_title(
        ws,
        "Legenda dei campi",
        (
            "Significato delle voci usate nel file. Nei fogli dati la stessa "
            "spiegazione compare anche come nota sull'intestazione della colonna."
        ),
        span=2,
    )
    add_back_link(ws, row=3)
    row = 5

    for section_name, pairs in blocks:
        row = section_title(ws, row, section_name, span=2)
        row = table_header(ws, row, ["Voce", "Significato"])
        for item_index, (name, meaning) in enumerate(pairs):
            name_cell = ws.cell(row=row, column=1, value=name)
            name_cell.font = LABEL_FONT
            name_cell.alignment = WRAP_LEFT
            name_cell.border = CELL_BORDER
            mean_cell = ws.cell(row=row, column=2, value=meaning)
            mean_cell.font = BASE_FONT
            mean_cell.alignment = WRAP_LEFT
            mean_cell.border = CELL_BORDER
            if item_index % 2:
                name_cell.fill = ZEBRA_FILL
                mean_cell.fill = ZEBRA_FILL
            row += 1
        row += 1

    ws.freeze_panes = "A5"
