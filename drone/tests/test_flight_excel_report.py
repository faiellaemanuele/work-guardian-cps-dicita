"""Test automatici relativi a flight excel report."""

from __future__ import annotations

import os
import sys
import tempfile
import textwrap
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from drone.data.flight_data_logger import FlightDataLogger

from flight_log_samples import _WP, _cmd, _pose, logger_with_autopilot_data, logger_with_both


def _openpyxl_available() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


# Nei fogli dati titolo, spiegazione e link al riepilogo occupano le prime
# righe: le intestazioni stanno nella riga 5 e i dati partono dalla 6.
_HEADER_ROW = 5
_FIRST_DATA_ROW = 6


def _fits_header_row(ws, header: str, width: float) -> bool:
    # L'intestazione va a capo nella larghezza della colonna: nessuna parola
    # deve sporgere e tutte le righe di testo devono stare nell'altezza della
    # riga delle intestazioni.
    from drone.data.excel_tables import HEADER_LINE_PT

    lines = textwrap.wrap(header, int(width), break_long_words=False)
    height = ws.row_dimensions[_HEADER_ROW].height
    return (all(len(line) <= width for line in lines)
            and len(lines) * HEADER_LINE_PT <= height)


def test_excel_has_expected_sheets_and_translations():
    if not _openpyxl_available():
        return
    from openpyxl import load_workbook
    from drone.data import flight_excel_report

    log = logger_with_both()
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "test.xlsx"
        flight_excel_report.save_session_workbook(log, path)
        wb = load_workbook(path)

    assert wb.sheetnames == ["Riepilogo", "Autopilota", "Kalman", "Legenda"]
    ws = wb["Autopilota"]
    headers = [ws.cell(_HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    assert "Cosa sta facendo il drone" in headers
    values = {ws.cell(r, headers.index("Cosa sta facendo il drone") + 1).value
              for r in range(_FIRST_DATA_ROW, ws.max_row + 1)}
    assert "Inseguimento waypoint" in values
    assert "tracking" not in values


def test_excel_columns_fit_headers():
    if not _openpyxl_available():
        return
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    from drone.data import flight_excel_report

    log = logger_with_both()
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "test.xlsx"
        flight_excel_report.save_session_workbook(log, path)
        wb = load_workbook(path)

    for sheet in ("Autopilota", "Kalman"):
        ws = wb[sheet]
        for c in range(1, ws.max_column + 1):
            header = ws.cell(_HEADER_ROW, c).value
            width = ws.column_dimensions[get_column_letter(c)].width
            assert width is not None and _fits_header_row(ws, header, width), (
                f"{sheet}: colonna '{header}' troppo stretta ({width})"
            )


def test_excel_parameters_sheet_from_config():
    if not _openpyxl_available():
        return
    from openpyxl import load_workbook
    from drone.config import APP_CONFIG
    from drone.data import flight_excel_report

    log = logger_with_autopilot_data()
    params = flight_excel_report.collect_run_parameters(log, APP_CONFIG, path_name="Percorso di prova")
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "test.xlsx"
        flight_excel_report.save_session_workbook(log, path, params)
        wb = load_workbook(path)

    assert "Parametri" in wb.sheetnames
    testo = "\n".join(
        str(wb["Parametri"].cell(r, c).value)
        for r in range(1, wb["Parametri"].max_row + 1)
        for c in (1, 2)
    )
    assert "Percorso di prova" in testo
    assert "Reattività negli spostamenti" in testo


def test_excel_parameters_accept_no_tag_distance_limit():
    from dataclasses import replace
    from drone.config import APP_CONFIG
    from drone.data import flight_excel_report

    config = replace(
        APP_CONFIG, camera_pose=replace(APP_CONFIG.camera_pose, max_tag_distance_m=None)
    )
    params = flight_excel_report.collect_run_parameters(FlightDataLogger(), config)
    voci = dict(voce for _gruppo, righe in params for voce in righe)
    assert voci["Distanza massima dei marker [m]"] == "nessun limite"


def _save_and_load(log, params=None):
    from openpyxl import load_workbook
    from drone.data import flight_excel_report

    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "test.xlsx"
        flight_excel_report.save_session_workbook(log, path, params)
        return load_workbook(path)


def _legend_entries(wb) -> set:
    ws = wb["Legenda"]
    return {
        ws.cell(r, 1).value for r in range(1, ws.max_row + 1)
        if ws.cell(r, 2).value is not None
    }


def _column_values(ws, header: str) -> set:
    headers = [ws.cell(_HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    colonna = headers.index(header) + 1
    return {ws.cell(r, colonna).value for r in range(_FIRST_DATA_ROW, ws.max_row + 1)}


def test_excel_explains_every_column_on_its_header():
    if not _openpyxl_available():
        return
    wb = _save_and_load(logger_with_both())
    for sheet in ("Autopilota", "Kalman"):
        ws = wb[sheet]
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(_HEADER_ROW, c)
            assert cell.comment is not None and cell.comment.text.strip(), (
                f"{sheet}: colonna '{cell.value}' senza spiegazione"
            )


def test_excel_legend_explains_every_parameter():
    if not _openpyxl_available():
        return
    from drone.config import APP_CONFIG
    from drone.data import flight_excel_report

    log = logger_with_autopilot_data()
    params = flight_excel_report.collect_run_parameters(log, APP_CONFIG, path_name="Percorso di prova")
    wb = _save_and_load(log, params)
    voci = _legend_entries(wb)
    ws = wb["Parametri"]
    etichette = [
        ws.cell(r, 1).value for r in range(1, ws.max_row + 1)
        if ws.cell(r, 2).value is not None
    ]
    assert etichette
    for etichetta in etichette:
        assert etichetta in voci, f"parametro '{etichetta}' assente dalla legenda"


# Tutti i motivi che l'autopilota può dare al comando e tutte le origini della
# posa che la stima e il filtro possono indicare.
_AUTOPILOT_REASONS = (
    "tracking", "z_priority_tracking", "waypoint_reached", "waypoint_timeout",
    "mission_finished", "pose_timeout", "pose_missing", "out_of_area",
    "supervision_stop", "supervision_stop_started", "supervision_stop_completed",
    "supervision_stop_tracking", "supervision_stop_z_priority_tracking",
    "invalid_pose: dati mancanti", "invalid_pose_transient: dati mancanti",
)
_POSE_SOURCES = ("weighted_average", "best_tag")


def test_excel_legend_explains_every_control_state_and_pose_source():
    if not _openpyxl_available():
        return
    log = FlightDataLogger()
    for i, reason in enumerate(_AUTOPILOT_REASONS):
        source = _POSE_SOURCES[i % len(_POSE_SOURCES)]
        filtrata = {**_pose(0.0, 0.0, 1.8, 0.0), "source": f"kalman({source})"}
        log.log_autopilot_step(
            pose_estimate=filtrata,
            command=_cmd(reason=reason),
            target=_WP(0.0, 1.0, 1.8, 0.0),
            timestamp=1000.0 + 0.05 * i,
        )
        log.log_pose_pair(
            raw_pose_estimate={**_pose(0.0, 0.0, 1.8, 0.0), "source": source},
            filtered_pose_estimate=filtrata,
            timestamp=1000.0 + 0.05 * i,
        )
    wb = _save_and_load(log)
    voci = _legend_entries(wb)

    stati = _column_values(wb["Autopilota"], "Cosa sta facendo il drone")
    assert len(stati) == len(_AUTOPILOT_REASONS)
    origini = (
        _column_values(wb["Autopilota"], "Metodo di stima della posizione")
        | _column_values(wb["Kalman"], "Metodo della posizione misurata")
        | _column_values(wb["Kalman"], "Metodo della posizione corretta")
    )
    for valore in stati | origini:
        assert valore in voci, f"valore '{valore}' assente dalla legenda"


def test_export_session_creates_excel():
    if not _openpyxl_available():
        return
    log = logger_with_autopilot_data()
    with tempfile.TemporaryDirectory() as d:
        session = log.export_session(output_root=d)
        assert session is not None
        assert (session / FlightDataLogger.EXCEL_FILENAME).exists()


def test_export_session_falls_back_to_text_without_excel():
    log = logger_with_autopilot_data()
    original = FlightDataLogger._save_excel
    FlightDataLogger._save_excel = lambda self, *a, **k: None
    try:
        with tempfile.TemporaryDirectory() as d:
            session = log.export_session(output_root=d)
            assert session is not None
            assert (session / FlightDataLogger.AUTOPILOT_TEXT_FILENAME).exists()
            assert not (session / FlightDataLogger.EXCEL_FILENAME).exists()
    finally:
        FlightDataLogger._save_excel = original


def _run_all() -> int:
    tests = sorted(
        (name, obj)
        for name, obj in globals().items()
        if name.startswith("test_") and callable(obj)
    )
    passed = 0
    failed = []
    for name, fn in tests:
        try:
            fn()
        except Exception as exc:  # noqa: BLE001
            failed.append((name, exc))
            print(f"[FAIL] {name}: {type(exc).__name__}: {exc}")
        else:
            passed += 1
            print(f"[ OK ] {name}")

    print("-" * 60)
    print(f"Totale: {len(tests)}  |  passati: {passed}  |  falliti: {len(failed)}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(_run_all())
