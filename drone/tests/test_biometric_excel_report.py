"""Test automatici relativi a biometric excel report."""

from __future__ import annotations

import os
import sys
import tempfile
import textwrap
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from drone.config import APP_CONFIG
from drone.data.biometric_data_logger import BiometricDataLogger

from biometric_log_samples import logger_con_sessione


def _openpyxl_available() -> bool:
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def _soglie():
    return APP_CONFIG.smartwatch_thresholds


def _workbook(log, *, con_parametri: bool = True):
    from openpyxl import load_workbook
    from drone.data import biometric_excel_report

    parametri = (
        biometric_excel_report.collect_session_parameters(log, APP_CONFIG)
        if con_parametri else None
    )
    with tempfile.TemporaryDirectory() as d:
        path = Path(d) / "test.xlsx"
        biometric_excel_report.save_session_workbook(log, path, _soglie(), parametri)
        return load_workbook(path)


# Nei fogli dati titolo, spiegazione e link al riepilogo occupano le prime
# righe: le intestazioni stanno nella riga 5 e i dati partono dalla 6.
_HEADER_ROW = 5
_FIRST_DATA_ROW = 6

# Nel riepilogo l'etichetta occupa le colonne da A a D e il valore la E.
_SUMMARY_VALUE_COL = 5


def _headers(ws):
    return [ws.cell(_HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]


def _fits_header_row(ws, header: str, width: float) -> bool:
    # L'intestazione va a capo nella larghezza della colonna: nessuna parola
    # deve sporgere e tutte le righe di testo devono stare nell'altezza della
    # riga delle intestazioni.
    from drone.data.excel_tables import HEADER_LINE_PT

    lines = textwrap.wrap(header, int(width), break_long_words=False)
    height = ws.row_dimensions[_HEADER_ROW].height
    return (all(len(line) <= width for line in lines)
            and len(lines) * HEADER_LINE_PT <= height)


def test_excel_has_a_sheet_for_the_watch_and_one_for_the_alarms():
    if not _openpyxl_available():
        return
    wb = _workbook(logger_con_sessione())
    assert wb.sheetnames == ["Riepilogo", "Parametri", "Orologio", "Allarmi", "Legenda"]


def test_the_watch_sheet_has_a_row_per_sample_and_translates_the_state():
    if not _openpyxl_available():
        return
    log = logger_con_sessione()
    ws = _workbook(log)["Orologio"]
    assert ws.max_row - _HEADER_ROW == len(log.samples())
    headers = _headers(ws)
    stati = {
        ws.cell(r, headers.index("Stato dell'orologio") + 1).value
        for r in range(_FIRST_DATA_ROW, ws.max_row + 1)
    }
    assert "Ricerca segnale" in stati
    assert "RICERCA_SEGNALE" not in stati


def test_the_watch_sheet_marks_the_samples_out_of_the_alarm_thresholds():
    if not _openpyxl_available():
        return
    ws = _workbook(logger_con_sessione())["Orologio"]
    headers = _headers(ws)
    fuori = [
        ws.cell(r, headers.index("Battito fuori soglia") + 1).value
        for r in range(_FIRST_DATA_ROW, ws.max_row + 1)
    ]
    # i campioni in allarme del campione di prova stanno a 132 bpm, oltre la
    # soglia di ingresso
    assert fuori.count("Sì") == 7
    assert fuori.count("No") == len(fuori) - 7


def test_the_watch_sheet_judges_the_thresholds_like_the_firmware():
    if not _openpyxl_available():
        return
    log = BiometricDataLogger()
    for i, hr in enumerate((120.4, 120.6)):
        log.log_message(
            {"bpm": round(hr), "spo2": 97, "stato": "NORMALE", "lettura_valida": True,
             "hr_filtrato": hr, "spo2_filtrato": 97.0},
            timestamp=100.0 + i * 0.5,
        )
    ws = _workbook(log)["Orologio"]
    headers = _headers(ws)
    colonna = headers.index("Battito fuori soglia") + 1
    # 120,4 bpm l'orologio li confronta come 120, dentro la soglia; 120,6 come 121
    righe = (_FIRST_DATA_ROW, _FIRST_DATA_ROW + 1)
    assert [ws.cell(r, colonna).value for r in righe] == ["No", "Sì"]


def test_the_alarm_time_is_counted_from_the_start_of_the_session():
    if not _openpyxl_available():
        return
    ws = _workbook(logger_con_sessione())["Allarmi"]
    headers = _headers(ws)
    relativo = ws.cell(_FIRST_DATA_ROW, headers.index("Tempo dall'avvio [s]") + 1).value
    # l'allarme del campione di prova cade dopo 25 campioni da mezzo secondo e
    # i dieci secondi di silenzio
    assert abs(relativo - (25 * 0.5 + 10.0)) < 1e-9
    assert ws.cell(_FIRST_DATA_ROW, headers.index("Causa") + 1).value == "Battito"


def test_without_alarms_the_workbook_has_no_alarm_sheet():
    if not _openpyxl_available():
        return
    log = BiometricDataLogger()
    for i in range(6):
        log.log_message(
            {
                "stato": "NORMALE", "lettura_valida": True,
                "hr_grezzo": 74.0, "hr_filtrato": 72.0,
                "spo2_grezzo": 98.0, "spo2_filtrato": 97.0,
            },
            timestamp=1000.0 + 0.5 * i,
        )
    wb = _workbook(log)
    assert "Allarmi" not in wb.sheetnames
    assert "Orologio" in wb.sheetnames


def test_the_alarm_sheet_has_no_watch_column():
    if not _openpyxl_available():
        return
    wb = _workbook(logger_con_sessione())
    assert _headers(wb["Allarmi"]) == [
        "Orario", "Tempo dall'avvio [s]", "Battito [bpm]", "Saturazione SpO₂ [%]", "Causa",
    ]


def test_the_columns_are_wide_enough_for_their_headers():
    if not _openpyxl_available():
        return
    from openpyxl.utils import get_column_letter

    wb = _workbook(logger_con_sessione())
    for sheet in ("Orologio", "Allarmi"):
        ws = wb[sheet]
        for c in range(1, ws.max_column + 1):
            header = ws.cell(_HEADER_ROW, c).value
            width = ws.column_dimensions[get_column_letter(c)].width
            assert width is not None and _fits_header_row(ws, header, width), (
                f"{sheet}: colonna '{header}' troppo stretta ({width})"
            )


def _legend_entries(wb) -> set:
    ws = wb["Legenda"]
    return {
        ws.cell(r, 1).value for r in range(1, ws.max_row + 1)
        if ws.cell(r, 2).value is not None
    }


def test_every_column_is_explained_on_its_header():
    if not _openpyxl_available():
        return
    wb = _workbook(logger_con_sessione())
    for sheet in ("Orologio", "Allarmi"):
        ws = wb[sheet]
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(_HEADER_ROW, c)
            assert cell.comment is not None and cell.comment.text.strip(), (
                f"{sheet}: colonna '{cell.value}' senza spiegazione"
            )


def test_the_legend_explains_every_parameter_and_watch_state():
    if not _openpyxl_available():
        return
    from drone.data.biometric_report_stats import STATE_LABELS

    wb = _workbook(logger_con_sessione())
    voci = _legend_entries(wb)
    ws = wb["Parametri"]
    etichette = [
        ws.cell(r, 1).value for r in range(1, ws.max_row + 1)
        if ws.cell(r, 2).value is not None
    ]
    assert etichette
    for etichetta in etichette:
        assert etichetta in voci, f"parametro '{etichetta}' assente dalla legenda"
    for stato in STATE_LABELS.values():
        assert stato in voci, f"stato '{stato}' assente dalla legenda"


def test_the_parameters_sheet_reports_the_firmware_thresholds():
    if not _openpyxl_available():
        return
    ws = _workbook(logger_con_sessione())["Parametri"]
    valori = {
        ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)
    }
    soglie = _soglie()
    assert valori["Allarme se il battito scende sotto [bpm]"] == str(soglie.bpm_min_in)
    assert valori["Allarme se il battito sale sopra [bpm]"] == str(soglie.bpm_max_in)
    assert (valori["Fine allarme quando la saturazione SpO₂ risale fino a [%]"]
            == str(soglie.spo2_min_out))


def test_the_summary_counts_the_telemetry_gap():
    if not _openpyxl_available():
        return
    ws = _workbook(logger_con_sessione())
    ws = ws["Riepilogo"]
    valori = {
        ws.cell(r, 1).value: ws.cell(r, _SUMMARY_VALUE_COL).value
        for r in range(1, ws.max_row + 1)
    }
    assert valori["Campioni biometrici ricevuti"] == 40
    assert valori["Interruzioni della telemetria [n.]"] == 1
    # i campioni coprono venti secondi: la sessione dura di più per il silenzio
    assert abs(valori["Durata con telemetria disponibile [s]"] - 20.0) < 1e-9
    assert valori["Durata della sessione [s]"] > 20.0
    assert abs(valori["Frequenza di acquisizione [campioni/s]"] - 2.0) < 1e-9


def test_export_session_writes_the_workbook():
    if not _openpyxl_available():
        return
    log = logger_con_sessione()
    with tempfile.TemporaryDirectory() as d:
        session = log.export_session(d, app_config=APP_CONFIG)
        assert session is not None
        assert (session / BiometricDataLogger.EXCEL_FILENAME).exists()
        assert not (session / BiometricDataLogger.TEXT_FILENAME).exists()


def test_export_session_falls_back_to_text_without_excel():
    log = logger_con_sessione()
    original = BiometricDataLogger._save_excel
    BiometricDataLogger._save_excel = lambda self, *a, **k: None
    try:
        with tempfile.TemporaryDirectory() as d:
            session = log.export_session(d, app_config=APP_CONFIG)
            assert session is not None
            assert (session / BiometricDataLogger.TEXT_FILENAME).exists()
            assert not (session / BiometricDataLogger.EXCEL_FILENAME).exists()
    finally:
        BiometricDataLogger._save_excel = original


def _run_all() -> int:
    tests = sorted(
        (name, obj)
        for name, obj in globals().items()
        if name.startswith("test_") and callable(obj)
    )
    passed = 0
    failed = []
    for name, fn in tests:
        try:
            fn()
        except Exception as exc:  # noqa: BLE001
            failed.append((name, exc))
            print(f"[FAIL] {name}: {type(exc).__name__}: {exc}")
        else:
            passed += 1
            print(f"[ OK ] {name}")

    print("-" * 60)
    print(f"Totale: {len(tests)}  |  passati: {passed}  |  falliti: {len(failed)}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(_run_all())
